"""Customer-facing data onboarding helpers.

This layer stages uploaded files, validates/matches columns, and delegates financial
persistence to the existing deterministic ingestion service. It never changes the
reconciliation engine or financial truth rules.
"""
from __future__ import annotations

import csv
import io
from datetime import datetime
from typing import Any

from app.ingestion import IngestionValidationError, ingest_records
from app.core.config import settings

MAX_UPLOAD_BYTES = settings.max_upload_bytes
MAX_ROWS = 10_000
MAX_PREVIEW_ROWS = 25
MAX_COLUMNS = 100
MAX_CELL_CHARS = 10_000

DATA_TYPES = {"payments", "settlements", "refunds", "bank_transactions"}

FIELD_DEFINITIONS: dict[str, dict[str, dict[str, Any]]] = {
    "payments": {
        "external_id": {"required": True, "aliases": ["external_id", "provider_payment_id", "payment_id", "transaction_id"]},
        "order_reference": {"required": False, "aliases": ["order_reference", "provider_order_id", "order_id", "merchant_ref", "merchant_reference"]},
        "amount": {"required": True, "aliases": ["amount", "payment_amount", "gross_amount", "total_amount"]},
        "currency": {"required": False, "aliases": ["currency", "currency_code"]},
        "status": {"required": False, "aliases": ["status", "payment_status"]},
        "event_timestamp": {"required": True, "aliases": ["event_timestamp", "created_at", "payment_date", "transaction_date", "date"]},
    },
    "settlements": {
        "external_id": {"required": True, "aliases": ["external_id", "provider_settlement_id", "settlement_id"]},
        "payment_reference": {"required": True, "aliases": ["payment_reference", "provider_payment_id", "payment_id", "payment_reference_id"]},
        "gross_amount": {"required": False, "aliases": ["gross_amount", "gross", "payment_amount", "amount"]},
        "fee": {"required": False, "aliases": ["fee", "fees", "gateway_fee", "processing_fee"]},
        "tax": {"required": False, "aliases": ["tax", "gst", "fee_tax"]},
        "net_amount": {"required": True, "aliases": ["net_amount", "net", "settlement_amount", "amount"]},
        "currency": {"required": False, "aliases": ["currency", "currency_code"]},
        "status": {"required": False, "aliases": ["status", "settlement_status"]},
        "event_timestamp": {"required": True, "aliases": ["event_timestamp", "settled_at", "settlement_date", "transaction_date", "date"]},
        "utr": {"required": False, "aliases": ["utr", "bank_reference", "utr_number"]},
    },
    "refunds": {
        "external_id": {"required": True, "aliases": ["external_id", "provider_refund_id", "refund_id"]},
        "payment_reference": {"required": True, "aliases": ["payment_reference", "provider_payment_id", "payment_id"]},
        "amount": {"required": True, "aliases": ["amount", "refund_amount"]},
        "currency": {"required": False, "aliases": ["currency", "currency_code"]},
        "status": {"required": False, "aliases": ["status", "refund_status"]},
        "event_timestamp": {"required": True, "aliases": ["event_timestamp", "created_at", "processed_at", "refund_date", "date"]},
    },
    "bank_transactions": {
        "external_id": {"required": True, "aliases": ["external_id", "transaction_id", "bank_transaction_id", "id"]},
        "reference": {"required": False, "aliases": ["reference", "description", "narration", "transaction_reference"]},
        "settlement_reference": {"required": False, "aliases": ["settlement_reference", "settlement_id"]},
        "utr": {"required": False, "aliases": ["utr", "utr_number", "bank_reference"]},
        "amount": {"required": True, "aliases": ["amount", "credit", "credit_amount", "transaction_amount"]},
        "currency": {"required": False, "aliases": ["currency", "currency_code"]},
        "status": {"required": False, "aliases": ["status", "transaction_status"]},
        "event_timestamp": {"required": True, "aliases": ["event_timestamp", "transaction_date", "value_date", "date", "created_at"]},
    },
}

DEFAULTS = {
    "currency": "INR",
    "status": {"payments": "captured", "settlements": "processed", "refunds": "processed", "bank_transactions": "posted"},
}


def _clean_header(value: Any) -> str:
    return str(value or "").strip().lower().replace(" ", "_").replace("-", "_")


def parse_upload(content: bytes, filename: str) -> tuple[list[dict[str, Any]], list[str]]:
    if not content:
        raise IngestionValidationError("The uploaded file is empty")
    if len(content) > MAX_UPLOAD_BYTES:
        raise IngestionValidationError("File exceeds the 5 MB upload limit")
    lower = filename.lower()
    if lower.endswith(".csv") or lower.endswith(".txt"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise IngestionValidationError("CSV must be UTF-8 encoded") from exc
        reader = csv.DictReader(io.StringIO(text))
        headers = [str(h).strip() for h in (reader.fieldnames or []) if h is not None]
        if not headers:
            raise IngestionValidationError("The CSV must contain a header row")
        if len(headers) > MAX_COLUMNS:
            raise IngestionValidationError(f"Files are limited to {MAX_COLUMNS} columns")
        rows = []
        for row in reader:
            if len(rows) >= MAX_ROWS:
                raise IngestionValidationError(f"Files are limited to {MAX_ROWS:,} rows")
            cleaned = {k.strip(): (v.strip() if isinstance(v, str) else v) for k, v in row.items() if k is not None}
            if any(len(v) > MAX_CELL_CHARS for v in cleaned.values() if isinstance(v, str)):
                raise IngestionValidationError(f"Each text cell is limited to {MAX_CELL_CHARS:,} characters")
            rows.append(cleaned)
        return rows, headers
    if lower.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise IngestionValidationError("Excel support is not installed on this server") from exc
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            iterator = sheet.iter_rows(values_only=True)
            header_values = next(iterator, None)
            headers = [str(v).strip() if v is not None else "" for v in (header_values or [])]
            headers = [h for h in headers if h]
            if not headers:
                raise IngestionValidationError("The Excel sheet must contain a header row")
            if len(headers) > MAX_COLUMNS:
                raise IngestionValidationError(f"Files are limited to {MAX_COLUMNS} columns")
            rows = []
            for values in iterator:
                if len(rows) >= MAX_ROWS:
                    raise IngestionValidationError(f"Files are limited to {MAX_ROWS:,} rows")
                row = {headers[i]: values[i] if i < len(values) else None for i in range(len(headers))}
                if any(len(v) > MAX_CELL_CHARS for v in row.values() if isinstance(v, str)):
                    raise IngestionValidationError(f"Each text cell is limited to {MAX_CELL_CHARS:,} characters")
                if any(v not in (None, "") for v in row.values()):
                    rows.append(row)
            workbook.close()
            return rows, headers
        except IngestionValidationError:
            raise
        except Exception as exc:
            raise IngestionValidationError("The Excel file could not be read") from exc
    raise IngestionValidationError("Unsupported file type. Upload CSV or XLSX")


def suggest_mapping(data_type: str, headers: list[str]) -> dict[str, str]:
    if data_type not in FIELD_DEFINITIONS:
        raise IngestionValidationError(f"Unsupported data type: {data_type}")
    normalized = {_clean_header(h): h for h in headers}
    mapping: dict[str, str] = {}
    for field, definition in FIELD_DEFINITIONS[data_type].items():
        for alias in definition["aliases"]:
            candidate = normalized.get(_clean_header(alias))
            if candidate:
                mapping[field] = candidate
                break
    return mapping


def validate_mapping(data_type: str, headers: list[str], mapping: dict[str, str]) -> list[str]:
    errors: list[str] = []
    available = set(headers)
    for field, definition in FIELD_DEFINITIONS[data_type].items():
        if definition["required"] and not mapping.get(field):
            errors.append(f"Missing mapping for required field: {field}")
        if mapping.get(field) and mapping[field] not in available:
            errors.append(f"Mapped column does not exist: {field}")
    return errors


def transform_rows(data_type: str, rows: list[dict[str, Any]], mapping: dict[str, str]) -> list[dict[str, Any]]:
    result = []
    for row_number, row in enumerate(rows, start=2):
        transformed: dict[str, Any] = {}
        for canonical, source_column in mapping.items():
            value = row.get(source_column)
            if isinstance(value, datetime):
                value = value.isoformat()
            transformed[canonical] = value
        transformed["currency"] = transformed.get("currency") or DEFAULTS["currency"]
        transformed["status"] = transformed.get("status") or DEFAULTS["status"][data_type]
        if data_type == "settlements":
            transformed["gross_amount"] = transformed.get("gross_amount") or transformed.get("net_amount")
            transformed["fee"] = transformed.get("fee") or "0"
            transformed["tax"] = transformed.get("tax") or "0"
        result.append(transformed)
    return result


def validate_rows(data_type: str, rows: list[dict[str, Any]], mapping: dict[str, str]) -> dict[str, Any]:
    mapping_errors = validate_mapping(data_type, list(mapping.values()), mapping)
    if mapping_errors:
        return {"valid": False, "rows": len(rows), "valid_rows": 0, "warning_rows": 0, "error_rows": len(rows), "errors": mapping_errors[:20], "warnings": []}
    transformed = transform_rows(data_type, rows, mapping)
    valid_rows = 0
    warning_rows = 0
    errors: list[str] = []
    warnings: list[str] = []
    for index, record in enumerate(transformed, start=2):
        try:
            # Reuse the authoritative normalizer for financial validation without persisting.
            from app.ingestion import normalize_record
            normalize_record(record, source="customer_upload", record_type={"payments": "payment", "settlements": "settlement", "refunds": "refund", "bank_transactions": "bank_transaction"}[data_type], provider="preview")
            valid_rows += 1
            if not record.get("order_reference") and data_type == "payments":
                warning_rows += 1
                if len(warnings) < 20: warnings.append(f"Row {index}: order reference is missing; matching may rely on payment ID or amount/date")
        except Exception as exc:
            if len(errors) < 20: errors.append(f"Row {index}: {str(exc)}")
    return {"valid": not errors, "rows": len(rows), "valid_rows": valid_rows, "warning_rows": warning_rows, "error_rows": len(rows) - valid_rows, "errors": errors, "warnings": warnings}


def import_rows(session, *, merchant_id: str, data_type: str, rows: list[dict[str, Any]], mapping: dict[str, str], source: str, provider: str | None) -> dict[str, Any]:
    transformed = transform_rows(data_type, rows, mapping)
    record_type = {"payments": "payment", "settlements": "settlement", "refunds": "refund", "bank_transactions": "bank_transaction"}[data_type]
    return ingest_records(session, transformed, source=source, record_type=record_type, provider=provider, merchant_id=merchant_id, commit=False)
